# -*- coding: utf-8 -*-
"""Escribe el papel de trabajo de revision de estados financieros en Excel.

    python generar_papel.py --estructura estructura.json --cuentas cuentas.json
        --ejercicios ejercicios.json --cliente "..." --ejercicio 2025
        --salida "<expediente>/InformesGesia/EstadosFinancieros/Estados Financieros <CLIENTE> <EJERCICIO>.xlsx"

Cinco hojas, y ninguna mas:

  Balance         la estructura oficial con los ejercicios cargados, variacion y % del ultimo
  PyG             igual, con el signo de presentacion de Gesia y el peso sobre ventas
  Conciliacion    SOLO el ejercicio corriente: por epigrafe, lo que presenta el cliente, los
                  ajustes, las reclasificaciones y lo auditado. Es el corazon de la revision
  Ajustes         el detalle por cuenta de lo anterior: quien mueve cada epigrafe
  Comprobaciones  los cuadres, con su semaforo. Si algo falla, se ve aqui y no hay que fiarse

**Sin autofiltro.** Lo pone el auditor si lo quiere, y donde lo quiera: puesto desde aqui se
ancla a la fila que el script cree que es la cabecera, y basta con que cambie el encabezado
para que aparezca una fila mas abajo, tapando la primera linea de datos con los desplegables.

**No usa formulas: los totales se calculan aqui, en Python, y se escriben como valor.** Un
papel de trabajo documenta lo que se midio el dia que se midio; si llevara formulas, cambiaria
de cifras al abrirlo seis meses despues y dejaria de ser evidencia. Es la regla de la casa, la
misma que en `cancelacion-saldos`.

Dos reglas de dominio que no son evidentes, medidas en la fase de contrato:

  * solo suman las cuentas con `AsignadaCCAA` y las lineas con `CuentasAsignables`; el plan
    repite la misma cuenta a varios niveles de digitos y las demas lineas son subtotales
  * el signo de presentacion NO es el del dato crudo: en el balance el pasivo viene con el
    signo cambiado, y en la PyG la regla es la de Gesia, epigrafes en magnitud y subtotales
    invertidos. Se replica tal cual para que las cifras coincidan con las que ve el auditor
"""

from __future__ import annotations

import argparse
import os
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from lib_eeff import (
    DECIMALES,
    ajustes_sin_epigrafe,
    ESTADO_BALANCE,
    ESTADO_PYG,
    TOLERANCIA,
    a_float,
    cuadre_balance,
    cuentas_por_epigrafe,
    redondear,
    ejercicios,
    es_si,
    estructura_ordenada,
    hojas,
    leer_json,
    salida_utf8,
    saldo,
    texto,
)

FONT = "Calibri"
CABECERA_FILL = PatternFill("solid", fgColor="1F4E78")
CABECERA_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")
AMARILLO = PatternFill("solid", fgColor="FFFF00")
GRIS = PatternFill("solid", fgColor="E7E6E6")
VERDE = Font(name=FONT, bold=True, color="006100", size=10)
ROJO = Font(name=FONT, bold=True, color="9C0006", size=10)
NORMAL = Font(name=FONT, size=10)
NEGRITA = Font(name=FONT, bold=True, size=10)
FINA = Side(style="thin", color="BFBFBF")
BORDE = Border(bottom=FINA)
EUROS = "#,##0.00"
PORCEN = "#,##0.0%"


def _cabecera(ws, fila: int, titulos: list, anchos: list) -> None:
    for i, (t, a) in enumerate(zip(titulos, anchos), start=1):
        c = ws.cell(row=fila, column=i, value=t)
        c.fill, c.font = CABECERA_FILL, CABECERA_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = a
    ws.freeze_panes = ws.cell(row=fila + 1, column=1)


def _titulo(ws, cliente: str, ejercicio: str, subtitulo: str) -> int:
    ws["A1"] = f"{cliente} · ejercicio {ejercicio}"
    ws["A1"].font = Font(name=FONT, bold=True, size=13)
    ws["A2"] = subtitulo
    ws["A2"].font = Font(name=FONT, size=10, italic=True, color="595959")
    ws["A3"] = ("Papel de trabajo preparado por Asistente IA el "
                + date.today().strftime("%d/%m/%Y")
                + ". Revisión: no formula las cuentas, las contrasta con el expediente.")
    ws["A3"].font = Font(name=FONT, size=9, color="808080")
    return 5


def _presentar_balance(fila: dict, indice: int) -> float:
    """El pasivo viene con el signo cambiado en la tabla cruda; se presenta en positivo."""
    v = saldo(fila, indice)
    return redondear(v if texto(fila.get("Activo_Pasivo")) == "ACTIVO" else -v)


def _presentar_pyg(fila: dict, indice: int) -> float:
    """La regla de Gesia, replicada: negativos a magnitud, y los subtotales invertidos.

    Se copia tal cual y no se 'mejora': el auditor compara este papel con la revision
    analitica del programa, y dos criterios de signo distintos son una discusion inutil.
    """
    v = saldo(fila, indice)
    clase = texto(fila.get("CodigoClase")).upper()
    if v < 0:
        return redondear(-v)
    if v > 0 and clase in ("ST", "T"):
        return redondear(-v)
    return redondear(v)


def _es_rotulo_pyg(fila: dict, estado: str) -> bool:
    """«A OPERACIONES CONTINUADAS» y «B OPERACIONES INTERRUMPIDAS»: rotulos de bloque.

    En el modelo del PGC son titulos sin cifra; las cifras estan en A.4) y en la linea 21. En
    la tabla de Gesia los dos rotulos guardan EL MISMO importe, el de las continuadas, medido
    igual en dos expedientes: el de interrumpidas es un reflejo, no un dato. Con la linea 21 a
    cero, escribirlo pondria varios millones de operaciones interrumpidas que no existen. Se
    escribe el rotulo y se deja la fila sin importes.
    """
    return estado == ESTADO_PYG and texto(fila.get("CodigoEpigrafe")) == "0"


def _sangria(fila: dict) -> int:
    niveles = sum(1 for c in ("CodigoSeccion", "CodigoEpigrafe", "CodigoDesglose", "CodigoSubDesglose")
                  if texto(fila.get(c)))
    return max(0, min(3, niveles - 1))


def _hoja_estado(wb: Workbook, nombre: str, estado: str, estructura: list, cols: list,
                 anios: dict, cliente: str, ejercicio: str, ventas: float = 0.0) -> None:
    ws = wb.create_sheet(nombre)
    sub = ("Balance de situación comparado, importes en euros."
           if estado == ESTADO_BALANCE else
           "Cuenta de pérdidas y ganancias comparada, importes en euros.")
    fila = _titulo(ws, cliente, ejercicio, sub)
    titulos = ["Epígrafe", "Concepto"] + [anios.get(str(i), f"col {i}") for i in cols]
    anchos = [14, 62] + [16] * len(cols)
    if len(cols) >= 2:
        titulos += ["Variación", "% var."]
        anchos += [16, 10]
    if estado == ESTADO_PYG and ventas:
        titulos += ["% s/ventas"]
        anchos += [11]
    _cabecera(ws, fila, titulos, anchos)
    fila += 1
    presentar = _presentar_balance if estado == ESTADO_BALANCE else _presentar_pyg
    for f in estructura_ordenada([x for x in estructura if texto(x.get("Estado")) == estado]):
        es_hoja = es_si(f.get("CuentasAsignables"))
        valores = [presentar(f, i) for i in cols]
        if es_hoja and not any(abs(v) > TOLERANCIA for v in valores):
            continue     # las lineas a cero en todos los ejercicios no aportan nada al papel
        ws.cell(row=fila, column=1, value=texto(f.get("CodigoAMostrar"))).font = NORMAL
        c = ws.cell(row=fila, column=2, value=texto(f.get("Concepto")))
        c.font = NORMAL if es_hoja else NEGRITA
        c.alignment = Alignment(indent=_sangria(f))
        if _es_rotulo_pyg(f, estado):
            for j in range(1, len(titulos) + 1):
                ws.cell(row=fila, column=j).fill = TOTAL_FILL
            ws.cell(row=fila, column=1).border = BORDE
            fila += 1
            continue
        for j, v in enumerate(valores, start=3):
            cel = ws.cell(row=fila, column=j, value=v)
            cel.number_format, cel.font = EUROS, (NORMAL if es_hoja else NEGRITA)
        col = 3 + len(cols)
        if len(cols) >= 2:
            var = redondear(valores[0] - valores[1])
            cel = ws.cell(row=fila, column=col, value=var)
            cel.number_format, cel.font = EUROS, (NORMAL if es_hoja else NEGRITA)
            pct = ws.cell(row=fila, column=col + 1,
                          value=(var / abs(valores[1])) if abs(valores[1]) > TOLERANCIA else None)
            pct.number_format, pct.font = PORCEN, (NORMAL if es_hoja else NEGRITA)
            col += 2
        if estado == ESTADO_PYG and ventas:
            cel = ws.cell(row=fila, column=col, value=(valores[0] / ventas) if ventas else None)
            cel.number_format, cel.font = PORCEN, (NORMAL if es_hoja else NEGRITA)
        if not es_hoja:
            for j in range(1, len(titulos) + 1):
                ws.cell(row=fila, column=j).fill = TOTAL_FILL
        ws.cell(row=fila, column=1).border = BORDE
        fila += 1


def _hoja_conciliacion(wb: Workbook, estructura: list, cuentas: list, cliente: str,
                       ejercicio: str) -> list:
    """Por epigrafe: lo presentado, lo ajustado, lo reclasificado y lo auditado.

    Devuelve las filas con diferencia, que son las que el auditor tiene que mirar.
    """
    ws = wb.create_sheet("Conciliacion")
    fila = _titulo(ws, cliente, ejercicio,
                   "Del saldo presentado por el cliente al saldo auditado, por epígrafe. "
                   "Solo el ejercicio corriente: el expediente no guarda el saldo de cliente de los anteriores.")
    _cabecera(ws, fila, ["Epígrafe", "Concepto", "Estado", "Según cliente", "Ajustes",
                         "Reclasificaciones", "Según auditoría", "Cuentas"],
              [14, 52, 10, 16, 15, 17, 16, 9])
    fila += 1
    cli = cuentas_por_epigrafe(cuentas, "SaldoCliente")
    aju = cuentas_por_epigrafe(cuentas, "SaldoAj")
    rec = cuentas_por_epigrafe(cuentas, "SaldoRec")
    aud = cuentas_por_epigrafe(cuentas, "SaldoAuditoria")
    cuantas: dict = {}
    for c in cuentas:
        if es_si(c.get("AsignadaCCAA")):
            cod = texto(c.get("CodigoCCAA"))
            cuantas[cod] = cuantas.get(cod, 0) + 1
    tocados = []
    for f in estructura_ordenada(hojas(estructura)):
        cod = texto(f.get("CodigoCCAA"))
        if cod not in aud:
            continue
        a, r = aju.get(cod, 0.0), rec.get(cod, 0.0)
        ws.cell(row=fila, column=1, value=texto(f.get("CodigoAMostrar"))).font = NORMAL
        ws.cell(row=fila, column=2, value=texto(f.get("Concepto"))).font = NORMAL
        ws.cell(row=fila, column=3, value=texto(f.get("Estado"))).font = NORMAL
        for j, v in enumerate((cli.get(cod, 0.0), a, r, aud.get(cod, 0.0)), start=4):
            cel = ws.cell(row=fila, column=j, value=redondear(v))
            cel.number_format, cel.font = EUROS, NORMAL
        ws.cell(row=fila, column=8, value=cuantas.get(cod, 0)).font = NORMAL
        if abs(a) > TOLERANCIA or abs(r) > TOLERANCIA:
            for j in (5, 6):
                ws.cell(row=fila, column=j).fill = AMARILLO
            tocados.append(texto(f.get("Concepto")))
        ws.cell(row=fila, column=1).border = BORDE
        fila += 1
    return tocados


def _hoja_ajustes(wb: Workbook, cuentas: list, estructura: list, cliente: str, ejercicio: str) -> int:
    ws = wb.create_sheet("Ajustes")
    fila = _titulo(ws, cliente, ejercicio,
                   "Las cuentas con ajuste o reclasificación de auditoría, en amarillo lo que mueve el saldo. "
                   "Solo las asignadas a un epígrafe: el plan repite la misma cuenta a 1, 2, 3 y 4 dígitos, "
                   "y listarlas todas enseñaría el mismo ajuste tres y cuatro veces.")
    _cabecera(ws, fila, ["Cuenta", "Nombre", "Epígrafe", "Según cliente", "Ajustes",
                         "Reclasificaciones", "Según auditoría"],
              [14, 46, 42, 16, 15, 17, 16])
    fila += 1
    concepto = {texto(f.get("CodigoCCAA")): texto(f.get("Concepto")) for f in estructura}
    n = 0
    for c in sorted(cuentas, key=lambda x: texto(x.get("Cuenta"))):
        # el ajuste de 4100 tambien esta en 410, en 41, en 4 y en el grupo 0 que Gesia usa para
        # cuadrar: la unica lectura que no multiplica el dinero es la de las cuentas asignadas
        if not es_si(c.get("AsignadaCCAA")):
            continue
        aj, rec = a_float(c.get("SaldoAj")), a_float(c.get("SaldoRec"))
        if abs(aj) <= TOLERANCIA and abs(rec) <= TOLERANCIA:
            continue
        n += 1
        ws.cell(row=fila, column=1, value=texto(c.get("Cuenta"))).font = NORMAL
        ws.cell(row=fila, column=2, value=texto(c.get("Nombre"))).font = NORMAL
        ws.cell(row=fila, column=3, value=concepto.get(texto(c.get("CodigoCCAA")), "")).font = NORMAL
        for j, v in enumerate((a_float(c.get("SaldoCliente")), aj, rec,
                               a_float(c.get("SaldoAuditoria"))), start=4):
            cel = ws.cell(row=fila, column=j, value=redondear(v))
            cel.number_format, cel.font = EUROS, NORMAL
            if j in (5, 6) and abs(v) > TOLERANCIA:
                cel.fill = AMARILLO
        ws.cell(row=fila, column=1).border = BORDE
        fila += 1
    if not n:
        ws.cell(row=fila, column=1, value="El expediente no tiene ajustes ni reclasificaciones.").font = NORMAL
    return n


def _hoja_comprobaciones(wb: Workbook, comprobaciones: list, cliente: str, ejercicio: str) -> None:
    ws = wb.create_sheet("Comprobaciones")
    fila = _titulo(ws, cliente, ejercicio,
                   "Los cuadres del papel. Si alguno falla, las cifras de las demás hojas no se sostienen.")
    todo = all(ok for _, ok, _ in comprobaciones)
    c = ws.cell(row=fila, column=1, value="TODO CUADRA" if todo else "HAY DESCUADRES: no entregar sin leerlos")
    c.font = VERDE if todo else ROJO
    fila += 2
    _cabecera(ws, fila, ["Comprobación", "Resultado", "Detalle"], [58, 14, 60])
    fila += 1
    for texto_c, ok, detalle in comprobaciones:
        ws.cell(row=fila, column=1, value=texto_c).font = NORMAL
        r = ws.cell(row=fila, column=2, value="CUADRA" if ok else "NO CUADRA")
        r.font = VERDE if ok else ROJO
        d = ws.cell(row=fila, column=3, value=detalle)
        d.font = NORMAL
        d.alignment = Alignment(wrap_text=True, vertical="top")
        if not ok:
            ws.cell(row=fila, column=1).fill = AMARILLO
        ws.cell(row=fila, column=1).border = BORDE
        fila += 1


def main() -> int:
    salida_utf8()
    p = argparse.ArgumentParser()
    p.add_argument("--estructura", required=True)
    p.add_argument("--cuentas", required=True)
    p.add_argument("--ejercicios", required=True)
    p.add_argument("--cliente", required=True)
    p.add_argument("--ejercicio", default="")
    p.add_argument("--salida", required=True)
    a = p.parse_args()

    estructura = leer_json(a.estructura)
    cuentas = leer_json(a.cuentas)
    anios = ejercicios(leer_json(a.ejercicios))

    cols = [i for i in range(1, 6) if any(abs(saldo(f, i)) > 0 for f in estructura)]
    if not cols:
        print("❌ ningun ejercicio tiene cifras: no se escribe nada")
        return 2
    ejercicio = a.ejercicio or anios.get(str(cols[0]), "")

    # ventas del ejercicio corriente, para el peso en la PyG: es el grupo 70 del plan
    ventas = 0.0
    for c in cuentas:
        if es_si(c.get("AsignadaCCAA")) and texto(c.get("Cuenta")).startswith("70"):
            ventas += abs(a_float(c.get("SaldoAuditoria")))
    ventas = round(ventas, DECIMALES)

    wb = Workbook()
    wb.remove(wb.active)
    _hoja_estado(wb, "Balance", ESTADO_BALANCE, estructura, cols, anios, a.cliente, ejercicio)
    if any(texto(f.get("Estado")) == ESTADO_PYG for f in estructura):
        _hoja_estado(wb, "PyG", ESTADO_PYG, estructura, cols, anios, a.cliente, ejercicio, ventas)
    tocados = _hoja_conciliacion(wb, estructura, cuentas, a.cliente, ejercicio)
    n_ajustes = _hoja_ajustes(wb, cuentas, estructura, a.cliente, ejercicio)

    comprobaciones = []
    for i in cols:
        d = cuadre_balance(estructura, i)
        comprobaciones.append((f"El balance de {anios.get(str(i), i)} cuadra: activo = patrimonio neto + pasivo",
                               abs(d) <= TOLERANCIA,
                               "diferencia 0,00" if abs(d) <= TOLERANCIA else f"difieren en {d:,.2f} €"))
    rotas = [texto(c.get("Cuenta")) for c in cuentas
             if abs((a_float(c.get("SaldoCliente")) + a_float(c.get("SaldoAj"))
                     + a_float(c.get("SaldoRec"))) - a_float(c.get("SaldoAuditoria"))) > TOLERANCIA]
    comprobaciones.append(("En cada cuenta, cliente + ajustes + reclasificaciones = auditoría",
                           not rotas,
                           f"cuadra en las {len(cuentas)} cuentas" if not rotas
                           else f"{len(rotas)} cuentas no cuadran: {', '.join(rotas[:5])}"))
    aud = cuentas_por_epigrafe(cuentas, "SaldoAuditoria")
    descuadres = []
    for f in hojas(estructura):
        cod = texto(f.get("CodigoCCAA"))
        if cod in aud:
            d = round(abs(saldo(f, cols[0])) - abs(aud[cod]), DECIMALES)
            if abs(d) > TOLERANCIA:
                descuadres.append(texto(f.get("Concepto")))
    comprobaciones.append(("Cada epígrafe es la suma de las cuentas asignadas a él",
                           not descuadres,
                           f"coincide en los {len(aud)} epígrafes con cuentas" if not descuadres
                           else f"{len(descuadres)} epígrafes no coinciden: {', '.join(descuadres[:3])}"))
    comprobaciones.append(("Las cuentas con ajuste o reclasificación están identificadas", True,
                           f"{n_ajustes} cuentas asignadas a un epígrafe, que mueven {len(tocados)} epígrafes"))
    fuera = ajustes_sin_epigrafe(cuentas)
    comprobaciones.append(("Ningún ajuste se queda fuera de los estados: toda cuenta ajustada "
                           "llega a un epígrafe, por sí misma o por su cuenta agregada",
                           not fuera,
                           "ninguna se queda fuera" if not fuera
                           else f"{len(fuera)} cuenta(s) con ajuste sin epígrafe: "
                                + "; ".join(f"{c} {n} ({v:,.2f} €)" for c, n, v in fuera[:4])))
    _hoja_comprobaciones(wb, comprobaciones, a.cliente, ejercicio)

    destino = os.path.abspath(a.salida)
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    print(f"✔ papel escrito: {destino}")
    print(f"  hojas: {', '.join(wb.sheetnames)}")
    print(f"  ejercicios: {', '.join(anios.get(str(i), str(i)) for i in cols)}")
    print(f"  cuentas con ajuste o reclasificacion: {n_ajustes}, en {len(tocados)} epigrafes")
    malas = [t for t, ok, _ in comprobaciones if not ok]
    if malas:
        print("\nPARA CONTAR AL ENTREGAR (tal cual, sin resumir):")
        for t, ok, d in comprobaciones:
            # el titulo es la AFIRMACION que se comprueba; si no se cumple, no se repite en
            # afirmativo pegada a su negacion (ver generar_efe.py, 25/09/2026)
            if not ok:
                print(f"  - NO CUADRA: {d}")
        return 1
    print("  todas las comprobaciones cuadran")
    return 0


if __name__ == "__main__":
    sys.exit(main())
