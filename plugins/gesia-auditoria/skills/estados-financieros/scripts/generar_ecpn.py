# -*- coding: utf-8 -*-
"""Escribe el papel de revision del patrimonio neto (ECPN) en Excel.

    python generar_ecpn.py --estructura estructura.json --cuentas cuentas.json
        --ejercicios ejercicios.json --cliente "..." --ejercicio 2025
        --salida "<expediente>/InformesGesia/EstadosFinancieros/Patrimonio Neto <CLIENTE> <EJERCICIO>.xlsx"

Cuatro hojas:

  PatrimonioNeto  los epigrafes de patrimonio neto de los ejercicios cargados, comparados
  Movimiento      EL ESTADO: por componente, de la apertura auditada al cierre auditado,
                  pasando por el movimiento del cliente, los ajustes y las reclasificaciones
  Detalle         lo mismo por cuenta: que cuenta mueve cada componente
  Comprobaciones  los cuadres, con su semaforo

**Gesia no tiene modulo de ECPN.** Se comprobo sacando del ejecutable los nombres de tabla
que aparecen en su propio SQL: estan los del balance, la PyG, la PyG analitica, las cuentas
anuales antiguas y las tres del estado de flujos, y ninguno de patrimonio neto. Asi que este
papel NO reproduce el ECPN oficial -no reparte el movimiento entre resultado, operaciones con
socios y otras variaciones, que es una clasificacion que hace el auditor-: **reconcilia cada
componente del patrimonio neto desde su apertura hasta su cierre**, que es lo que se puede
sostener con el dato que hay, y deja el reparto a quien firma.

La clasificacion por componentes NO se inventa aqui: es la del propio Gesia, la del epigrafe
del balance al que esta asignada cada cuenta. Capital, prima, reservas, resultados anteriores,
resultado del ejercicio, otras aportaciones, ajustes por cambios de valor y subvenciones salen
de ahi, no de leer el numero de cuenta.

Como en el resto del skill, se escriben VALORES y no formulas.
"""

from __future__ import annotations

import argparse
import os
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from lib_eeff import (
    TOLERANCIA,
    a_float,
    ejercicios,
    es_si,
    estructura_ordenada,
    hojas,
    leer_json,
    lineas_pn,
    movimiento_cuenta,
    movimiento_por_epigrafe,
    redondear,
    salida_utf8,
    saldo,
    texto,
)

FONT = "Calibri"
CABECERA_FILL = PatternFill("solid", fgColor="1F4E78")
CABECERA_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")
AMARILLO = PatternFill("solid", fgColor="FFFF00")
VERDE = Font(name=FONT, bold=True, color="006100", size=10)
ROJO = Font(name=FONT, bold=True, color="9C0006", size=10)
GRIS_FONT = Font(name=FONT, bold=True, color="7F6000", size=10)
NORMAL = Font(name=FONT, size=10)
NEGRITA = Font(name=FONT, bold=True, size=10)
EUROS = "#,##0.00"

# Grupos del PGC para los ingresos y gastos imputados directamente al patrimonio neto. Si el
# cliente no los usa -los dos expedientes de calibracion no los usan-, el estado de ingresos y
# gastos reconocidos no se puede desglosar desde la contabilidad y hay que decirlo.
GRUPOS_IMPUTACION = ("8", "9")


def _pn(v: float) -> float:
    """El patrimonio neto viene acreedor, con el signo contrario al del activo.

    Se presenta en positivo, igual que en el balance del papel de estados financieros, para que
    un patrimonio neto sano se lea positivo. Uno negativo sigue saliendo negativo: eso no se
    maquilla, que es justo lo que el auditor tiene que ver.
    """
    return redondear(-v)


def _cabecera(ws, fila: int, titulos: list, anchos: list) -> None:
    for i, (t, a) in enumerate(zip(titulos, anchos), start=1):
        c = ws.cell(row=fila, column=i, value=t)
        c.fill, c.font = CABECERA_FILL, CABECERA_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = a
    ws.freeze_panes = ws.cell(row=fila + 1, column=1)


def _titulo(ws, cliente: str, ejercicio: str, subtitulo: str) -> int:
    ws["A1"] = f"{cliente} · ejercicio {ejercicio}"
    ws["A1"].font = Font(name=FONT, bold=True, size=13)
    ws["A2"] = subtitulo
    ws["A2"].font = Font(name=FONT, size=10, italic=True, color="595959")
    ws["A3"] = ("Papel de trabajo preparado por Asistente IA el "
                + date.today().strftime("%d/%m/%Y")
                + ". Revisión: reconcilia el patrimonio neto del expediente, no formula el ECPN.")
    ws["A3"].font = Font(name=FONT, size=9, color="808080")
    return 5


def _importe(ws, fila: int, col: int, valor: float, negrita: bool = False, marcar: bool = False):
    c = ws.cell(row=fila, column=col, value=redondear(valor))
    c.number_format, c.font = EUROS, (NEGRITA if negrita else NORMAL)
    if marcar and abs(valor) > TOLERANCIA:
        c.fill = AMARILLO
    return c


def _hoja_comparado(wb: Workbook, pn: list, cols: list, anios: dict, cliente: str,
                    ejercicio: str) -> None:
    ws = wb.create_sheet("PatrimonioNeto")
    fila = _titulo(ws, cliente, ejercicio,
                   "Los epígrafes de patrimonio neto de los ejercicios cargados, en euros y en "
                   "signo de presentación.")
    titulos = ["Epígrafe", "Concepto"] + [anios.get(str(i), f"col {i}") for i in cols]
    anchos = [14, 56] + [17] * len(cols)
    if len(cols) >= 2:
        titulos.append("Variación")
        anchos.append(17)
    _cabecera(ws, fila, titulos, anchos)
    fila += 1
    for f in estructura_ordenada(pn):
        es_hoja = es_si(f.get("CuentasAsignables"))
        valores = [_pn(saldo(f, i)) for i in cols]
        if es_hoja and not any(abs(v) > TOLERANCIA for v in valores):
            continue
        ws.cell(row=fila, column=1, value=texto(f.get("CodigoAMostrar"))).font = NORMAL
        c = ws.cell(row=fila, column=2, value=texto(f.get("Concepto")))
        c.font = NORMAL if es_hoja else NEGRITA
        for j, v in enumerate(valores, start=3):
            _importe(ws, fila, j, v, negrita=not es_hoja)
        if len(cols) >= 2:
            _importe(ws, fila, 3 + len(cols), valores[0] - valores[1], negrita=not es_hoja)
        if not es_hoja:
            for j in range(1, len(titulos) + 1):
                ws.cell(row=fila, column=j).fill = TOTAL_FILL
        fila += 1


def _hoja_movimiento(wb: Workbook, pn: list, cuentas: list, cliente: str,
                     ejercicio: str, anterior: str) -> tuple:
    """EL estado. Devuelve (totales, descuadres por epigrafe)."""
    ws = wb.create_sheet("Movimiento")
    fila = _titulo(ws, cliente, ejercicio,
                   "De la apertura al cierre, componente a componente. La apertura es el saldo "
                   f"AUDITADO de {anterior or 'el ejercicio anterior'}, no el del cliente.")
    _cabecera(ws, fila,
              ["Epígrafe", "Componente", f"Apertura ({anterior or 'N-1'}) auditada",
               "Movimiento del cliente", "Ajustes de auditoría", "Reclasificaciones",
               "Cierre auditado"],
              [14, 46, 20, 20, 18, 18, 20])
    fila += 1
    codigos = {texto(f.get("CodigoCCAA")) for f in hojas(pn)}
    mov = movimiento_por_epigrafe(cuentas, codigos)
    totales = [0.0] * 5
    descuadres = []
    for f in estructura_ordenada(hojas(pn)):
        cod = texto(f.get("CodigoCCAA"))
        if cod not in mov:
            continue
        ant, m, aj, rec, cie = (_pn(v) for v in mov[cod])
        if not any(abs(v) > TOLERANCIA for v in (ant, m, aj, rec, cie)):
            continue
        ws.cell(row=fila, column=1, value=texto(f.get("CodigoAMostrar"))).font = NORMAL
        ws.cell(row=fila, column=2, value=texto(f.get("Concepto"))).font = NORMAL
        _importe(ws, fila, 3, ant)
        _importe(ws, fila, 4, m)
        _importe(ws, fila, 5, aj, marcar=True)
        _importe(ws, fila, 6, rec, marcar=True)
        _importe(ws, fila, 7, cie)
        if abs((ant + m + aj + rec) - cie) > TOLERANCIA:
            descuadres.append((texto(f.get("Concepto")), redondear(ant + m + aj + rec - cie)))
            ws.cell(row=fila, column=7).fill = AMARILLO
        for i, v in enumerate((ant, m, aj, rec, cie)):
            totales[i] = redondear(totales[i] + v)
        fila += 1
    ws.cell(row=fila, column=2, value="TOTAL PATRIMONIO NETO").font = NEGRITA
    for j, v in enumerate(totales, start=3):
        _importe(ws, fila, j, v, negrita=True)
    for j in range(1, 8):
        ws.cell(row=fila, column=j).fill = TOTAL_FILL
    fila += 2
    ws.cell(row=fila, column=2,
            value="La variación del ejercicio es " f"{redondear(totales[4] - totales[0]):,.2f} €: "
                  "el movimiento del cliente más los ajustes y reclasificaciones de auditoría. "
                  "Repartirla entre resultado, operaciones con socios y otras variaciones es "
                  "del auditor, no de este papel.").font = Font(name=FONT, size=9, italic=True,
                                                                color="595959")
    return totales, descuadres


def _hoja_detalle(wb: Workbook, pn: list, cuentas: list, cliente: str, ejercicio: str) -> int:
    ws = wb.create_sheet("Detalle")
    fila = _titulo(ws, cliente, ejercicio,
                   "Qué cuenta mueve cada componente. Solo las cuentas asignadas a un epígrafe: "
                   "el plan repite la misma cuenta a 1, 2, 3 y 4 dígitos.")
    _cabecera(ws, fila,
              ["Cuenta", "Nombre", "Componente", "Apertura", "Movimiento", "Ajustes",
               "Reclasificaciones", "Cierre"],
              [12, 42, 34, 18, 18, 16, 17, 18])
    fila += 1
    codigos = {texto(f.get("CodigoCCAA")) for f in hojas(pn)}
    concepto = {texto(f.get("CodigoCCAA")): texto(f.get("Concepto")) for f in pn}
    n = 0
    for c in sorted(cuentas, key=lambda x: texto(x.get("Cuenta"))):
        cod = texto(c.get("CodigoCCAA"))
        if cod not in codigos or not es_si(c.get("AsignadaCCAA")):
            continue
        valores = [_pn(v) for v in movimiento_cuenta(c)]
        if not any(abs(v) > TOLERANCIA for v in valores):
            continue
        n += 1
        ws.cell(row=fila, column=1, value=texto(c.get("Cuenta"))).font = NORMAL
        ws.cell(row=fila, column=2, value=texto(c.get("Nombre"))).font = NORMAL
        ws.cell(row=fila, column=3, value=concepto.get(cod, "")).font = NORMAL
        for j, v in enumerate(valores, start=4):
            _importe(ws, fila, j, v, marcar=(j in (6, 7)))
        fila += 1
    if not n:
        ws.cell(row=fila, column=1,
                value="Ninguna cuenta asignada a patrimonio neto tiene saldo ni movimiento.").font = NORMAL
    return n


def _hoja_comprobaciones(wb: Workbook, comprobaciones: list, cliente: str, ejercicio: str) -> None:
    ws = wb.create_sheet("Comprobaciones")
    fila = _titulo(ws, cliente, ejercicio,
                   "Los cuadres del papel. Si alguno falla, las cifras de las demás hojas no se "
                   "sostienen.")
    todo = all(ok is not False for _, ok, _ in comprobaciones)
    c = ws.cell(row=fila, column=1,
                value="TODO CUADRA" if todo else "HAY DESCUADRES: no entregar sin leerlos")
    c.font = VERDE if todo else ROJO
    fila += 2
    _cabecera(ws, fila, ["Comprobación", "Resultado", "Detalle"], [62, 14, 62])
    fila += 1
    for t, ok, detalle in comprobaciones:
        # tres estados, no dos: hay cosas que no son un cuadre y no pueden salir en verde.
        # Un «CUADRA» al lado de «el cliente no usa los grupos 8 y 9» se lee como que todo
        # esta bien, y lo que dice es que falta informacion para el estado de ingresos y
        # gastos reconocidos.
        ws.cell(row=fila, column=1, value=t).font = NORMAL
        etiqueta = "A TENER EN CUENTA" if ok is None else ("CUADRA" if ok else "NO CUADRA")
        r = ws.cell(row=fila, column=2, value=etiqueta)
        r.font = GRIS_FONT if ok is None else (VERDE if ok else ROJO)
        d = ws.cell(row=fila, column=3, value=detalle)
        d.font = NORMAL
        d.alignment = Alignment(wrap_text=True, vertical="top")
        if ok is False:
            ws.cell(row=fila, column=1).fill = AMARILLO
        fila += 1


def main() -> int:
    salida_utf8()
    p = argparse.ArgumentParser()
    p.add_argument("--estructura", required=True)
    p.add_argument("--cuentas", required=True)
    p.add_argument("--ejercicios", required=True)
    p.add_argument("--cliente", required=True)
    p.add_argument("--ejercicio", default="")
    p.add_argument("--salida", required=True)
    a = p.parse_args()

    estructura = leer_json(a.estructura)
    cuentas = leer_json(a.cuentas)
    anios = ejercicios(leer_json(a.ejercicios))

    if not cuentas or "SaldoAnterior" not in cuentas[0]:
        print("❌ el plan exportado no trae SaldoAnterior: sin la apertura no hay estado de "
              "cambios que valga. Vuelve a exportar las cuentas incluyendo c.SaldoAnterior.")
        return 2

    pn = lineas_pn(estructura)
    if not pn:
        print("❌ no hay lineas de patrimonio neto en la estructura del balance")
        return 2

    cols = [i for i in range(1, 6) if any(abs(saldo(f, i)) > 0 for f in estructura)]
    if not cols:
        print("❌ ningun ejercicio tiene cifras: no se escribe nada")
        return 2
    ejercicio = a.ejercicio or anios.get(str(cols[0]), "")
    anterior = anios.get(str(cols[1]), "") if len(cols) >= 2 else ""

    wb = Workbook()
    wb.remove(wb.active)
    _hoja_comparado(wb, pn, cols, anios, a.cliente, ejercicio)
    totales, descuadres = _hoja_movimiento(wb, pn, cuentas, a.cliente, ejercicio, anterior)
    n_cuentas = _hoja_detalle(wb, pn, cuentas, a.cliente, ejercicio)

    comprobaciones = []
    comprobaciones.append((
        "En cada componente, apertura + movimiento + ajustes + reclasificaciones = cierre",
        not descuadres,
        f"cuadra en todos los componentes con saldo" if not descuadres
        else f"{len(descuadres)} no cuadran: "
             + "; ".join(f"{c} ({d:,.2f} €)" for c, d in descuadres[:3])))

    # el cierre del estado contra el epigrafe TOTAL del balance del ejercicio corriente
    total_balance = 0.0
    for f in hojas(pn):
        total_balance = redondear(total_balance + _pn(saldo(f, cols[0])))
    d = redondear(totales[4] - total_balance)
    comprobaciones.append((
        "El cierre del estado es el patrimonio neto del balance del ejercicio corriente",
        abs(d) <= TOLERANCIA,
        f"{total_balance:,.2f} €" if abs(d) <= TOLERANCIA else f"difieren en {d:,.2f} €"))

    # LA CONTINUIDAD: la apertura contra el patrimonio neto auditado del ejercicio anterior
    if len(cols) >= 2:
        total_anterior = 0.0
        for f in hojas(pn):
            total_anterior = redondear(total_anterior + _pn(saldo(f, cols[1])))
        d = redondear(totales[0] - total_anterior)
        comprobaciones.append((
            f"La apertura es el patrimonio neto auditado de {anterior or 'N-1'}, sin saltos",
            abs(d) <= TOLERANCIA,
            f"{total_anterior:,.2f} €" if abs(d) <= TOLERANCIA
            else f"difieren en {d:,.2f} €: el ejercicio abre con un patrimonio neto distinto "
                 "del que cerró el anterior, y eso hay que explicarlo"))

    # el resultado del ejercicio del estado contra el de la cuenta de resultados
    resultado_pyg = None
    for f in estructura:
        if texto(f.get("Estado")) == "PyG" and texto(f.get("CodigoAMostrar")).startswith("A.5"):
            resultado_pyg = redondear(-saldo(f, cols[0]))
    if resultado_pyg is not None:
        codigos = {texto(f.get("CodigoCCAA")) for f in hojas(pn)}
        mov = movimiento_por_epigrafe(cuentas, codigos)
        # por el CONCEPTO, no por el numero de epigrafe: el «7.» del modelo normal no es el
        # mismo epigrafe en el abreviado ni en el de pymes, y cuadrar contra la linea
        # equivocada da un verde que no vale nada
        lineas_res = [f for f in hojas(pn)
                      if "resultado del ejercicio" in texto(f.get("Concepto")).lower()]
        if not lineas_res:
            comprobaciones.append((
                "El resultado del ejercicio del patrimonio neto es el de la cuenta de resultados",
                None,
                "no se ha encontrado la línea de resultado del ejercicio en el patrimonio neto "
                "de este modelo de balance: la comprobación no se ha podido hacer"))
        else:
            resultado_pn = 0.0
            for f in lineas_res:
                resultado_pn = redondear(resultado_pn + _pn(mov.get(texto(f.get("CodigoCCAA")), [0] * 5)[4]))
            d = redondear(resultado_pn - resultado_pyg)
            comprobaciones.append((
                "El resultado del ejercicio del patrimonio neto es el de la cuenta de resultados",
                abs(d) <= TOLERANCIA,
                f"{resultado_pyg:,.2f} €" if abs(d) <= TOLERANCIA else f"difieren en {d:,.2f} €"))

    # los grupos 8 y 9: si no se usan, el estado de ingresos y gastos reconocidos no se desglosa
    usa_89 = any(texto(c.get("Cuenta"))[:1] in GRUPOS_IMPUTACION
                 and abs(a_float(c.get("SaldoAuditoria"))) > TOLERANCIA for c in cuentas)
    comprobaciones.append((
        "Los grupos 8 y 9 (ingresos y gastos imputados al patrimonio neto) están contabilizados",
        True if usa_89 else None,
        "sí: el estado de ingresos y gastos reconocidos se puede desglosar desde la contabilidad"
        if usa_89 else
        "NO los usa el cliente. El patrimonio neto cuadra igual, pero el estado de ingresos y "
        "gastos reconocidos no sale de la contabilidad: hay que construirlo con el auditor"))

    _hoja_comprobaciones(wb, comprobaciones, a.cliente, ejercicio)

    destino = os.path.abspath(a.salida)
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    wb.save(destino)
    print(f"✔ papel escrito: {destino}")
    print(f"  hojas: {', '.join(wb.sheetnames)}")
    print(f"  patrimonio neto: apertura {totales[0]:,.2f} → cierre {totales[4]:,.2f} "
          f"(variación {redondear(totales[4] - totales[0]):,.2f})")
    print(f"  cuentas de patrimonio neto con saldo o movimiento: {n_cuentas}")
    malas = [t for t, ok, _ in comprobaciones if ok is False]
    avisos = [t for t, ok, _ in comprobaciones if ok is None]
    if malas or avisos:
        print("\nPARA CONTAR AL ENTREGAR (tal cual, sin resumir):")
        for t, ok, d in comprobaciones:
            if ok is False:
                print(f"  - NO CUADRA: {d}")
            elif ok is None:
                print(f"  - A TENER EN CUENTA: {d}")
        return 1
    print("  todas las comprobaciones cuadran")
    return 0


if __name__ == "__main__":
    sys.exit(main())
