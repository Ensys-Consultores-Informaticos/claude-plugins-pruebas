# -*- coding: utf-8 -*-
"""Arnes de estados-financieros. Comprobaciones PURAS: no toca ningun expediente.

    python probar_eeff.py

Lo que se prueba aqui es lo que se rompe en silencio: los signos de presentacion, que solo
sumen las lineas y las cuentas que deben, y que el orden sea el del modelo oficial y no el
alfabetico. Un fallo en cualquiera de esos tres da un papel con pinta correcta y cifras falsas.
"""

from __future__ import annotations

import io
import os
import sys

from lib_eeff import (
    ESTADO_BALANCE,
    lineas_pn,
    movimiento_cuenta,
    movimiento_por_epigrafe,
    ajustes_sin_epigrafe,
    ESTADO_PYG,
    a_float,
    cadena_del_saldo,
    cuadre_balance,
    cuentas_por_epigrafe,
    ejercicios,
    es_si,
    estructura_ordenada,
    hojas,
    salida_utf8,
    saldo,
)
from generar_ecpn import _pn
from generar_efe import ESTRUCTURALES, _agrupar, _fecha, _neto, _totales
from proponer_efe import (
    PREFIJO_IA,
    Saldos,
    _hoja_catalogo,
    agrupar,
    proponer,
    regla_de,
    total_propuesto,
)
from generar_papel import _es_rotulo_pyg, _presentar_balance, _presentar_pyg, _sangria


def _catalogo_de_prueba():
    """La hoja Catalogo de un expediente con dos plantillas: una con propuesta y otra sin."""
    from openpyxl import Workbook

    catalogo = [
        {"NumeroAsientoOA": "3", "Descripcion": "Amortizacion", "Aprobado": "False",
         "Observaciones": "", "ApunteOA": str(i), "Cuenta": cta, "Contrapartida": con,
         "Concepto": "x", "Origen": "0", "Aplicacion": "0"}
        for i, (cta, con) in enumerate(
            [("C68", "280"), ("280", "C68"), ("C68", "281"), ("281", "C68"),
             ("C68", "282"), ("282", "C68")], start=1)
    ] + [
        {"NumeroAsientoOA": "9", "Descripcion": "Lo que rellena el auditor", "Aprobado": "False",
         "Observaciones": "", "ApunteOA": "9", "Cuenta": "570", "Contrapartida": "129",
         "Concepto": "x", "Origen": "0", "Aplicacion": "0"},
    ]
    grupos = agrupar(catalogo)
    props = proponer(grupos, Saldos([{"Cuenta": "680", "SaldoAuditoria": "414,32"},
                                     {"Cuenta": "681", "SaldoAuditoria": "103.896,35"}]))
    wb = Workbook()
    wb.remove(wb.active)
    _hoja_catalogo(wb, grupos, {p["numero"]: total_propuesto(p) for p in props},
                   "CLIENTE DE PRUEBA", "2025")
    return wb["Catalogo"]


def _linea(codigo, estado, ap, clase, asignable, s1=0.0, s2=0.0, **extra):
    f = {"CodigoCCAA": codigo, "Estado": estado, "Activo_Pasivo": ap, "CodigoClase": clase,
         "CuentasAsignables": asignable, "SaldoAuditoria1": s1, "SaldoAuditoria2": s2,
         "Concepto": "x", "CodigoAMostrar": ""}
    f.update(extra)
    return f


def main() -> int:
    salida_utf8()
    fallos = 0

    # un balance de juguete: activo 1.000, pasivo -1.000 en crudo, mas un subtotal que
    # NO debe sumar (CuentasAsignables False)
    balance = [
        _linea("Activo A)", ESTADO_BALANCE, "ACTIVO", "A", "False", 1000.0, 900.0),
        _linea("Activo A)I.", ESTADO_BALANCE, "ACTIVO", "A", "True", 600.0, 500.0),
        _linea("Activo A)II.", ESTADO_BALANCE, "ACTIVO", "A", "True", 400.0, 400.0),
        _linea("Pasivo A)I.", ESTADO_BALANCE, "PASIVO", "P", "True", -1000.0, -900.0),
    ]
    pyg = [
        _linea("I.", ESTADO_PYG, "", "I", "True", -500.0, -400.0),      # ingreso, en crudo negativo
        _linea("II.", ESTADO_PYG, "", "I", "True", 120.0, 100.0),       # gasto, en crudo positivo
        _linea("A)", ESTADO_PYG, "", "ST", "False", 380.0, 300.0),      # subtotal positivo
    ]
    cuentas = [
        {"Cuenta": "7000", "CodigoCCAA": "I.", "AsignadaCCAA": "True", "SaldoAuditoria": -500.0,
         "SaldoCliente": -450.0, "SaldoAj": -50.0, "SaldoRec": 0.0, "Nombre": "Ventas"},
        {"Cuenta": "700", "CodigoCCAA": "I.", "AsignadaCCAA": "False", "SaldoAuditoria": -500.0,
         "SaldoCliente": -450.0, "SaldoAj": -50.0, "SaldoRec": 0.0, "Nombre": "agregada, NO suma"},
        {"Cuenta": "6000", "CodigoCCAA": "II.", "AsignadaCCAA": "True", "SaldoAuditoria": 120.0,
         "SaldoCliente": 120.0, "SaldoAj": 0.0, "SaldoRec": 0.0, "Nombre": "Compras"},
    ]

    pruebas = [
        ("importes: coma decimal española, punto de millar, y el formato ingles",
         a_float("1.234,56") == 1234.56 and a_float("1234.56") == 1234.56
         and a_float("") == 0.0 and a_float(None) == 0.0 and a_float("-1.000,00") == -1000.0),

        ("los booleanos del API llegan como texto",
         es_si("True") and es_si("-1") and es_si("Sí") and not es_si("False") and not es_si("")),

        ("solo son hojas las lineas que admiten cuentas: el subtotal queda fuera",
         len(hojas(balance)) == 3 and len(hojas(pyg)) == 2),

        ("EL CUADRE: activo mas pasivo suma cero sobre las hojas, en los dos ejercicios",
         abs(cuadre_balance(balance, 1)) < 0.01 and abs(cuadre_balance(balance, 2)) < 0.01),

        ("un balance descuadrado NO se da por bueno",
         abs(cuadre_balance(balance + [_linea("Pasivo A)II.", ESTADO_BALANCE, "PASIVO", "P", "True", -5.0)], 1)) > 0.01),

        ("presentacion del balance: el activo tal cual y el pasivo cambiado de signo",
         _presentar_balance(balance[1], 1) == 600.0 and _presentar_balance(balance[3], 1) == 1000.0),

        ("presentacion de la PyG, regla de Gesia: el negativo a magnitud, el positivo de epigrafe igual",
         _presentar_pyg(pyg[0], 1) == 500.0 and _presentar_pyg(pyg[1], 1) == 120.0),

        ("presentacion de la PyG: el subtotal POSITIVO se invierte",
         _presentar_pyg(pyg[2], 1) == -380.0),

        ("agregacion por epigrafe: solo cuentan las asignadas, la agregada no duplica",
         cuentas_por_epigrafe(cuentas, "SaldoAuditoria") == {"I.": -500.0, "II.": 120.0}),

        ("la cadena del saldo: cliente + ajustes + reclasificaciones = auditoria",
         (lambda c: abs((c[0] + c[1] + c[2]) - c[3]) < 0.01)(cadena_del_saldo(cuentas[0]))),

        ("los ejercicios se leen de Auditorias y el 1 no se supone que sea un año",
         ejercicios([{"CodigoAuditoria": "1", "Año": "2025"}, {"CodigoAuditoria": "2", "Año": "2024"}])
         == {"1": "2025", "2": "2024"}),

        ("orden del modelo, no alfabetico: el epigrafe 2 va antes que el 10",
         [f["CodigoCCAA"] for f in estructura_ordenada([
             _linea("b", ESTADO_BALANCE, "ACTIVO", "A", "True", CodigoSeccion="A", CodigoEpigrafe="10"),
             _linea("a", ESTADO_BALANCE, "ACTIVO", "A", "True", CodigoSeccion="A", CodigoEpigrafe="2")])]
         == ["a", "b"]),

        ("LAS SECCIONES NO SE MEZCLAN: el no corriente entero antes que el corriente, aunque "
         "la seccion sea una LETRA y el epigrafe del corriente sea menor",
         [f["CodigoCCAA"] for f in estructura_ordenada([
             _linea("B.1", ESTADO_BALANCE, "ACTIVO", "A", "True", CodigoSeccion="B", CodigoEpigrafe="1"),
             _linea("A.2", ESTADO_BALANCE, "ACTIVO", "A", "True", CodigoSeccion="A", CodigoEpigrafe="2"),
             _linea("A.hd", ESTADO_BALANCE, "ACTIVO", "A", "False", CodigoSeccion="A")])]
         == ["A.hd", "A.2", "B.1"]),

        ("la PyG se ordena por CodigoOrdenI, y dentro de un mismo orden el detalle va antes "
         "que el subtotal: «20. Impuestos» antes de «A.4) RESULTADO...»",
         [f["CodigoCCAA"] for f in estructura_ordenada([
             _linea("A.4)", ESTADO_PYG, "", "T", "False", CodigoOrdenI="20", CodigoSeccion="4"),
             _linea("21.", ESTADO_PYG, "", "I", "True", CodigoOrdenI="21", CodigoEpigrafe="21"),
             _linea("20.", ESTADO_PYG, "", "I", "True", CodigoOrdenI="20", CodigoEpigrafe="20")])]
         == ["20.", "A.4)", "21."]),

        ("ningun ajuste se pierde: el de 4100 llega por 410, el grupo 0 de cuadre de Gesia no "
         "cuenta, y una cuenta ajustada sin epigrafe ni agregada SI se denuncia",
         [c for c, _, _ in ajustes_sin_epigrafe([
             {"Cuenta": "4100", "AsignadaCCAA": "False", "SaldoAj": -53.07, "SaldoRec": 0.0, "Nombre": ""},
             {"Cuenta": "410", "AsignadaCCAA": "True", "SaldoAj": -53.07, "SaldoRec": 0.0, "Nombre": ""},
             {"Cuenta": "080", "AsignadaCCAA": "False", "SaldoAj": 999.0, "SaldoRec": 0.0, "Nombre": "Explotacion"},
             {"Cuenta": "8990", "AsignadaCCAA": "False", "SaldoAj": 12.0, "SaldoRec": 0.0, "Nombre": "huerfana"},
         ])] == ["8990"]),

        ("el balance va antes que la PyG, y dentro del balance el activo antes que el pasivo",
         [f["Estado"] for f in estructura_ordenada(balance + pyg)][:4] == [ESTADO_BALANCE] * 4),

        ("los rotulos de bloque de la PyG van SIN cifra: «B OPERACIONES INTERRUMPIDAS» guarda "
         "el importe de las continuadas y escribirlo seria inventar un resultado",
         _es_rotulo_pyg(_linea("B", ESTADO_PYG, "", "I", "False", 5432881.57, CodigoEpigrafe="0"), ESTADO_PYG)
         and not _es_rotulo_pyg(_linea("1.", ESTADO_PYG, "", "I", "True", 100.0, CodigoEpigrafe="1"), ESTADO_PYG)
         and not _es_rotulo_pyg(_linea("A)", ESTADO_BALANCE, "ACTIVO", "A", "False", 100.0, CodigoEpigrafe="0"), ESTADO_BALANCE)),

        ("la sangria crece con el nivel del epigrafe, y no se pasa de tres",
         _sangria(_linea("x", ESTADO_BALANCE, "ACTIVO", "A", "True", CodigoSeccion="A")) == 0
         and _sangria(_linea("x", ESTADO_BALANCE, "ACTIVO", "A", "True", CodigoSeccion="A",
                             CodigoEpigrafe="1", CodigoDesglose="1", CodigoSubDesglose="1")) == 3),

        ("el patrimonio neto son las lineas del PASIVO de la seccion A, y se reconoce por ahi "
         "y no por el texto del rotulo, que cambia entre modelos",
         [f["CodigoCCAA"] for f in lineas_pn([
             _linea("PN", ESTADO_BALANCE, "PASIVO", "P", "True", CodigoSeccion="A"),
             _linea("exigible", ESTADO_BALANCE, "PASIVO", "P", "True", CodigoSeccion="B"),
             _linea("activo", ESTADO_BALANCE, "ACTIVO", "A", "True", CodigoSeccion="A"),
             _linea("pyg", ESTADO_PYG, "", "I", "True", CodigoSeccion="A")])] == ["PN"]),

        ("EL MOVIMIENTO DEL EJERCICIO es cierre menos apertura, NO el Debe menos el Haber: "
         "esos acumulados llevan dentro el saldo de apertura",
         movimiento_cuenta({"SaldoAnterior": "-2.328.453,48", "DebeCliente": "0",
                            "HaberCliente": "2.876.858,37", "SaldoCliente": "-2.876.858,37",
                            "SaldoAj": "0", "SaldoRec": "0",
                            "SaldoAuditoria": "-2.876.858,37"})[1] == -548404.89),

        ("apertura + movimiento + ajustes + reclasificaciones = cierre, y solo suman las "
         "cuentas asignadas al componente",
         (lambda m: abs(sum(m["Pasivo A)I.3."][:4]) - m["Pasivo A)I.3."][4]) < 0.01
                 and abs(m["Pasivo A)I.3."][0] - (-2328453.48)) < 0.01)(
             movimiento_por_epigrafe([
                 {"Cuenta": "113", "CodigoCCAA": "Pasivo A)I.3.", "AsignadaCCAA": "True",
                  "SaldoAnterior": -2328453.48, "SaldoCliente": -2876858.37, "SaldoAj": 145327.30,
                  "SaldoRec": 0.0, "SaldoAuditoria": -2731531.07},
                 {"Cuenta": "11", "CodigoCCAA": "Pasivo A)I.3.", "AsignadaCCAA": "False",
                  "SaldoAnterior": -2328453.48, "SaldoCliente": -2876858.37, "SaldoAj": 145327.30,
                  "SaldoRec": 0.0, "SaldoAuditoria": -2731531.07},
             ], {"Pasivo A)I.3."}))),

        ("el patrimonio neto se presenta en positivo, y uno NEGATIVO sigue saliendo negativo",
         _pn(-61404.0) == 61404.0 and _pn(867.04) == -867.04),

        ("EFE · el ajuste es partida doble ENTRE LINEAS DEL ESTADO, y sus cobros tienen que "
         "igualar a sus pagos: un descuadre se ve",
         _totales([{"Origen": "100,00", "Aplicacion": "0"},
                   {"Origen": "0", "Aplicacion": "100,00"}])[2] == 0.0
         and _totales([{"Origen": "100,00", "Aplicacion": "0"}])[2] == 100.0),

        ("EFE · los ajustes llegan de un LEFT JOIN, con la cabecera repetida en cada apunte y "
         "la plantilla sin usar sin ninguno",
         (lambda g: len(g) == 2 and len(g["1"]["apuntes"]) == 2 and g["1"]["aprobado"]
                 and g["7"]["apuntes"] == [] and not g["7"]["aprobado"])(
             _agrupar([
                 {"NumeroAsientoOA": "1", "Descripcion": "Amortizacion", "Aprobado": "True",
                  "ApunteOA": "1", "Origen": "10", "Aplicacion": "0"},
                 {"NumeroAsientoOA": "1", "Descripcion": "Amortizacion", "Aprobado": "True",
                  "ApunteOA": "2", "Origen": "0", "Aplicacion": "10"},
                 {"NumeroAsientoOA": "7", "Descripcion": "Plantilla sin usar", "Aprobado": "False",
                  "ApunteOA": "", "Origen": "", "Aplicacion": ""},
             ]))),

        ("EFE · el estado se lee en neto, cobro menos pago",
         _neto({"OrigenAjustado": "1.541.971,78", "AplicacionAjustada": "0"},
               "OrigenAjustado", "AplicacionAjustada") == 1541971.78
         and _neto({"OrigenAjustado": "0", "AplicacionAjustada": "92.933,78"},
                   "OrigenAjustado", "AplicacionAjustada") == -92933.78),

        ("EFE · las lineas estructurales del cuadro se reconocen por CodigoLinea, que en las de "
         "detalle lleva la cuenta o el concepto",
         all(t in ESTRUCTURALES for t in ("F", "T", "S", "P"))
         and not any(t in ESTRUCTURALES for t in ("129", "C68", "G630", "57"))),

        ("EFE · la fecha viene con la hora pegada y cortarla a diez deja una fecha que no existe",
         _fecha("31/12/2025 0:00:00") == "31/12/2025" and _fecha("31/12/2025") == "31/12/2025"
         and _fecha(None) == ""),

        ("PROPUESTA · la plantilla se reconoce por LA FORMA de sus apuntes, no por el texto de "
         "la descripcion, que el auditor puede cambiar",
         regla_de({"apuntes": [{"Cuenta": "C68", "Contrapartida": "280"},
                               {"Cuenta": "C68", "Contrapartida": "281"},
                               {"Cuenta": "C68", "Contrapartida": "282"}]})[0]
         == "Amortización del ejercicio"
         and regla_de({"apuntes": [{"Cuenta": "999", "Contrapartida": "998"}]})[0] is None),

        ("PROPUESTA · los saldos se leen distinto segun el grupo: la de resultados por su SALDO "
         "del año, la de balance por su MOVIMIENTO",
         (lambda s: s.importe("66") == 220875.62 and s.importe("113") == 548404.89)(
             Saldos([{"Cuenta": "66", "SaldoAuditoria": "220.875,62", "SaldoAnterior": "0",
                      "SaldoCliente": "74.629,14"},
                     {"Cuenta": "113", "SaldoAuditoria": "-2.731.531,07",
                      "SaldoAnterior": "-2.328.453,48", "SaldoCliente": "-2.876.858,37"}]))),

        ("PROPUESTA · una cuenta que NO esta en el plan da None y no cero: proponer cero seria "
         "inventarse que el concepto no se ha movido",
         Saldos([]).importe("680") is None and Saldos([]).apertura("129") is None),

        ("PROPUESTA · con el resultado anterior en PERDIDA la regla de beneficios no propone "
         "nada, que es lo correcto mientras la de perdidas no se mida",
         proponer(agrupar([
             {"NumeroAsientoOA": "1", "Descripcion": "Distribucion", "Aprobado": "False",
              "Observaciones": "", "ApunteOA": "1", "Cuenta": "129", "Contrapartida": "120",
              "Concepto": "(C)", "Origen": "0", "Aplicacion": "0"},
             {"NumeroAsientoOA": "1", "Descripcion": "Distribucion", "Aprobado": "False",
              "Observaciones": "", "ApunteOA": "2", "Cuenta": "11", "Contrapartida": "129",
              "Concepto": "(P)", "Origen": "0", "Aplicacion": "0"},
             {"NumeroAsientoOA": "1", "Descripcion": "Distribucion", "Aprobado": "False",
              "Observaciones": "", "ApunteOA": "3", "Cuenta": "526", "Contrapartida": "129",
              "Concepto": "(P)", "Origen": "0", "Aplicacion": "0"},
         ]), Saldos([{"Cuenta": "129", "SaldoAnterior": "5.176.303,90",
                      "SaldoCliente": "5.432.881,57", "SaldoAuditoria": "5.432.881,57"}])) == []),

        ("PROPUESTA · la amortizacion empareja cada dotacion con SU acumulada, el tramo sin "
         "dotacion se queda A CERO en vez de desaparecer, y la propuesta cuadra",
         (lambda ps: len(ps) == 1
                 and dict(((l, c), (co, pa)) for (l, c), _, co, pa in ps[0]["lineas"])
                     == {("C68", "280"): (414.32, 0.0), ("280", "C68"): (0.0, 414.32),
                         ("C68", "281"): (103896.35, 0.0), ("281", "C68"): (0.0, 103896.35),
                         ("C68", "282"): (0.0, 0.0), ("282", "C68"): (0.0, 0.0)}
                 and abs(sum(co for _, _, co, _ in ps[0]["lineas"])
                         - sum(pa for _, _, _, pa in ps[0]["lineas"])) < 0.01)(
             proponer(agrupar([
                 {"NumeroAsientoOA": "3", "Descripcion": "Amortizacion", "Aprobado": "False",
                  "Observaciones": "", "ApunteOA": str(i), "Cuenta": cta, "Contrapartida": con,
                  "Concepto": "x", "Origen": "0", "Aplicacion": "0"}
                 for i, (cta, con) in enumerate(
                     [("C68", "280"), ("280", "C68"), ("C68", "281"), ("281", "C68"),
                      ("C68", "282"), ("282", "C68")], start=1)
             ]), Saldos([{"Cuenta": "680", "SaldoAuditoria": "414,32"},
                         {"Cuenta": "681", "SaldoAuditoria": "103.896,35"}])))),

        ("PROPUESTA · lo que se propone se firma: el prefijo «Asistente IA:» va en la "
         "descripcion, para que no se confunda nunca con el trabajo del auditor",
         PREFIJO_IA.startswith("Asistente IA:")),

        ("ENTREGA · la linea que se cuenta al auditor NO repite el titulo de la comprobacion: "
         "el titulo es la afirmacion que se comprueba, y pegado a su detalle se leia al reves "
         "-«El estado esta trabajado ... NINGUNO esta aprobado»- (25/09/2026)",
         all('{t} \u00b7 {d}' not in io.open(
                 os.path.join(os.path.dirname(os.path.abspath(__file__)), g), encoding="utf-8").read()
             for g in ("generar_efe.py", "generar_papel.py", "generar_ecpn.py", "proponer_efe.py"))),

        ("ENTREGA · y el aviso de que no hay ningun ajuste aprobado sigue diciendo que el cuadro "
         "cuadra consigo mismo aunque el estado no este hecho, que es lo que enganya",
         (lambda t: "no es" in t and "todav" in t and "verde de las dem" in t)(
             io.open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "generar_efe.py"),
                     encoding="utf-8").read())),

        ("CATALOGO · el importe propuesto se ve en la hoja Catalogo, al lado del ya aprobado y "
         "con otro nombre. Los dos compartian la columna «Importe puesto», que suma lo YA "
         "REGISTRADO en el expediente -0,00 mientras nada esta aprobado-, y se leyo como que la "
         "propuesta venia vacia aunque la hoja Propuestas traia las cifras (26/09/2026)",
         (lambda ws: [ws.cell(row=5, column=c).value for c in (5, 6)]
          == ["Importe ya aprobado en Gesia", "Importe propuesto"]
          and ws.cell(row=6, column=5).value == 0.0
          and ws.cell(row=6, column=6).value == 104310.67)(_catalogo_de_prueba())),

        ("CATALOGO · y la plantilla SIN propuesta deja la celda en blanco, no a 0,00: un cero "
         "ahi es justo la lectura que se quiere evitar",
         _catalogo_de_prueba().cell(row=7, column=6).value is None),

        ("REDONDEO CONTABLE: el medio centimo sube, y no se redondea al par como hace round()",
         saldo({"SaldoAuditoria1": "1.000,555"}, 1) == 1000.56
         and saldo({"SaldoAuditoria1": "0,125"}, 1) == 0.13
         and saldo({"SaldoAuditoria1": "0,135"}, 1) == 0.14
         and round(0.125, 2) == 0.12),
    ]

    for descripcion, ok in pruebas:
        print(f"  {'OK   ' if ok else 'FALLA'}  {descripcion}")
        if not ok:
            fallos += 1

    print()
    print("RESULTADO:", "todo correcto" if fallos == 0 else f"{fallos} fallo(s)")
    return 1 if fallos else 0


if __name__ == "__main__":
    sys.exit(main())
